"""Build manuscript-ready Word tables for the fitted mammalian parameters.

The numerical values are read directly from the current saved fit outputs.
The document uses the compact_reference_guide preset, with one named visual
override: Times New Roman and quiet three-rule tables to match the supplied
supplementary-table reference.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
FIT_DIR = PROJECT_DIR / "Output" / "scipy_sensor_logR2_floor_0p5"
LBD_CSV = FIT_DIR / "parameters_lbd.csv"
DBD_CSV = FIT_DIR / "parameters_dbd.csv"
README = PROJECT_DIR / "README.md"
OUTPUT_DOCX = FIT_DIR / "mammalian_fitted_DBD_LBD_parameters.docx"

FONT = "Times New Roman"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


LBD_ORDER = [
    "RpaR179",
    "BjaR180",
    "TraR174",
    "CinR179",
    "CarRecc169",
    "BjaR180S107R",
    "LasR177",
    "acVHH",
    "ER282-595",
    "DHBR282-595",
    "PR",
    "MR669-984",
    "GR487-777",
]

INDUCER_BY_LBD = {
    "RpaR179": "pC-HSL",
    "BjaR180": "IV-HSL",
    "TraR174": "3OC8-HSL",
    "CinR179": "3OHC14-HSL",
    "CarRecc169": "3OC8-HSL",
    "BjaR180S107R": "IV-HSL",
    "LasR177": "3OC12-HSL",
    "acVHH": "Caffeine",
    "ER282-595": "β-estradiol",
    "DHBR282-595": "DHB",
    "PR": "RU486",
    "MR669-984": "Aldosterone",
    "GR487-777": "Dexamethasone",
}

NAME_PARTS = {
    "RpaR179": [("RpaR", "normal"), ("179", "sub")],
    "BjaR180": [("BjaR", "normal"), ("180", "sub")],
    "TraR174": [("TraR", "normal"), ("174", "sub")],
    "CinR179": [("CinR", "normal"), ("179", "sub")],
    "CarRecc169": [("CarRecc", "normal"), ("169", "sub")],
    "BjaR180S107R": [
        ("BjaR", "normal"),
        ("180", "sub"),
        ("S107R", "super"),
    ],
    "LasR177": [("LasR", "normal"), ("177", "sub")],
    "acVHH": [("acVHH", "normal")],
    "ER282-595": [("ER", "normal"), ("282-595", "sub")],
    "DHBR282-595": [("DHBR", "normal"), ("282-595", "sub")],
    "PR": [("PR", "normal")],
    "MR669-984": [("MR", "normal"), ("669-984", "sub")],
    "GR487-777": [("GR", "normal"), ("487-777", "sub")],
    "CI43470": [("CI434", "normal"), ("70", "sub")],
    "lexAec87": [("lexA", "normal"), ("ec87", "sub")],
    "lexAxa101": [("lexA", "normal"), ("xa101", "sub")],
    "PurR60": [("PurR", "normal"), ("60", "sub")],
    "lexAmm115": [("lexA", "normal"), ("mm115", "sub")],
    "lexAfs104": [("lexA", "normal"), ("fs104", "sub")],
    "CI94": [("CI", "normal"), ("94", "sub")],
    "HKCI84": [("HKCI", "normal"), ("84", "sub")],
    "lexAgs91": [("lexA", "normal"), ("gs91", "sub")],
    "RecApact104": [("RecApact", "normal"), ("104", "sub")],
    "DeoR92": [("DeoR", "normal"), ("92", "sub")],
}


def read_csv_indexed(path: Path, key: str) -> tuple[list[str], dict[str, dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    order = [row[key] for row in rows]
    if len(set(order)) != len(order):
        raise ValueError(f"Duplicate {key} rows in {path}")
    return order, {row[key]: row for row in rows}


def strip_md(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value.startswith("`") and value.endswith("`"):
        value = value[1:-1]
    return value.strip()


def read_dbd_aliases(path: Path) -> dict[str, list[str]]:
    aliases: defaultdict[str, list[str]] = defaultdict(list)
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        if not raw_line.lstrip().startswith("|"):
            continue
        cells = [strip_md(cell) for cell in raw_line.strip().strip("|").split("|")]
        if len(cells) < 6 or not re.fullmatch(r"\d+\.csv", cells[0]):
            continue
        source_dbd, canonical_dbd, fit_status = cells[2], cells[3], cells[5]
        if fit_status.casefold() != "included":
            continue
        if source_dbd not in aliases[canonical_dbd]:
            aliases[canonical_dbd].append(source_dbd)
    return dict(aliases)


def set_run_font(run, *, size: float = 11, bold: bool = False, italic: bool = False):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def style_cell_paragraph(paragraph, alignment: WD_ALIGN_PARAGRAPH) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def add_name(paragraph, name: str, *, bold: bool = False) -> None:
    clear_paragraph(paragraph)
    for text, position in NAME_PARTS.get(name, [(name, "normal")]):
        run = set_run_font(paragraph.add_run(text), size=10.5, bold=bold)
        if position == "sub":
            run.font.subscript = True
        elif position == "super":
            run.font.superscript = True


def add_parameter_symbol(paragraph, base: str, subscript: str) -> None:
    clear_paragraph(paragraph)
    run = set_run_font(paragraph.add_run(base), size=10.5, bold=True, italic=True)
    run = set_run_font(paragraph.add_run(subscript), size=10.5, bold=True, italic=True)
    run.font.subscript = True


def add_scientific_value(paragraph, value: float, *, fixed_marker: bool = False) -> None:
    clear_paragraph(paragraph)
    value = float(value)
    magnitude = abs(value)
    if magnitude != 0 and (magnitude < 1e-3 or magnitude >= 1e3):
        exponent = int(math.floor(math.log10(magnitude)))
        coefficient = value / (10**exponent)
        set_run_font(paragraph.add_run(f"{coefficient:.2f}×10"), size=10.5)
        exponent_run = set_run_font(paragraph.add_run(str(exponent)), size=10.5)
        exponent_run.font.superscript = True
    else:
        set_run_font(paragraph.add_run(format(value, ".3g")), size=10.5)
    if fixed_marker:
        marker = set_run_font(paragraph.add_run("a"), size=8)
        marker.font.superscript = True


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS_DXA.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Column widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_quiet_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        if edge in {"top", "bottom"}:
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "10")
            node.set(qn("w:color"), "000000")
            node.set(qn("w:space"), "0")
        else:
            node.set(qn("w:val"), "nil")

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        bottom.set(qn("w:space"), "0")


def add_caption(document: Document, table_number: int, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = 1.0
    run = set_run_font(
        paragraph.add_run(f"Supplementary Table {table_number}: {title}"),
        size=11.5,
        bold=True,
    )
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_note(document: Document, parts: list[tuple]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for part in parts:
        text, italic = part[:2]
        superscript = bool(part[2]) if len(part) > 2 else False
        run = set_run_font(paragraph.add_run(text), size=9, italic=italic)
        run.font.superscript = superscript
        run.font.color.rgb = RGBColor(55, 55, 55)


def build_lbd_table(document: Document, lbd_rows: dict[str, dict[str, str]]) -> None:
    add_caption(document, 4, "Full-curve fitted apparent LBD parameters in mammalian cells.")
    table = document.add_table(rows=1, cols=5)
    set_table_geometry(table, [1980, 1900, 1820, 1820, 1840])
    set_quiet_table_borders(table)
    set_repeat_table_header(table.rows[0])

    headers = table.rows[0].cells
    for paragraph in [cell.paragraphs[0] for cell in headers]:
        style_cell_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    clear_paragraph(headers[0].paragraphs[0])
    set_run_font(headers[0].paragraphs[0].add_run("LBD"), size=10.5, bold=True)
    clear_paragraph(headers[1].paragraphs[0])
    set_run_font(headers[1].paragraphs[0].add_run("Inducer"), size=10.5, bold=True)
    add_parameter_symbol(headers[2].paragraphs[0], "K", "d0")
    add_parameter_symbol(headers[3].paragraphs[0], "K", "b")
    add_parameter_symbol(headers[4].paragraphs[0], "K", "d1")

    for name in LBD_ORDER:
        source = lbd_rows[name]
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
        style_cell_paragraph(cells[0].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
        add_name(cells[0].paragraphs[0], name)
        style_cell_paragraph(cells[1].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
        clear_paragraph(cells[1].paragraphs[0])
        set_run_font(cells[1].paragraphs[0].add_run(INDUCER_BY_LBD[name]), size=10.5)
        for cell, field in zip(cells[2:], ("Kd0", "Kb_uM_inverse", "Kd1")):
            style_cell_paragraph(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
            add_scientific_value(cell.paragraphs[0], float(source[field]))

    set_table_geometry(table, [1980, 1900, 1820, 1820, 1840])
    add_note(
        document,
        [
            ("Note. ", True),
            ("Values are rounded to three significant figures. Units: Kb, µM⁻¹; Kd0 and Kd1, RPU⁻¹. Parameters are shared by all sensors containing the same canonical LBD. Q_LBD = Kb²Kd1/Kd0 is derived and is not listed as an independently fitted parameter.", False),
        ],
    )


def build_dbd_table(
    document: Document,
    dbd_order: list[str],
    dbd_rows: dict[str, dict[str, str]],
    dbd_aliases: dict[str, list[str]],
) -> None:
    document.add_page_break()
    add_caption(document, 5, "Full-curve fitted apparent DBD parameters in mammalian cells.")
    table = document.add_table(rows=1, cols=4)
    set_table_geometry(table, [2150, 3050, 2000, 2160])
    set_quiet_table_borders(table)
    set_repeat_table_header(table.rows[0])

    headers = table.rows[0].cells
    for paragraph in [cell.paragraphs[0] for cell in headers]:
        style_cell_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    clear_paragraph(headers[0].paragraphs[0])
    set_run_font(headers[0].paragraphs[0].add_run("DBD"), size=10.5, bold=True)
    clear_paragraph(headers[1].paragraphs[0])
    set_run_font(headers[1].paragraphs[0].add_run("Source-data DBD name(s)"), size=10.5, bold=True)
    add_parameter_symbol(headers[2].paragraphs[0], "K", "A")
    add_parameter_symbol(headers[3].paragraphs[0], "T", "0,variant")

    for name in dbd_order:
        source = dbd_rows[name]
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
        style_cell_paragraph(cells[0].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
        add_name(cells[0].paragraphs[0], name)
        style_cell_paragraph(cells[1].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
        clear_paragraph(cells[1].paragraphs[0])
        set_run_font(cells[1].paragraphs[0].add_run("; ".join(dbd_aliases.get(name, []))), size=10.5)
        style_cell_paragraph(cells[2].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        add_scientific_value(
            cells[2].paragraphs[0],
            float(source["KA"]),
            fixed_marker=source["KA_fixed"].casefold() == "true",
        )
        style_cell_paragraph(cells[3].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        add_scientific_value(cells[3].paragraphs[0], float(source["T0_variant"]))

    set_table_geometry(table, [2150, 3050, 2000, 2160])
    add_note(
        document,
        [
            ("Note. ", True),
            ("Values are rounded to three significant figures. Units: KA, RPU⁻¹; T0,variant, RPU. The same canonical DBD shares KA and T0,variant across sensors. ", False),
            ("a", False, True),
            (" KA for lexAec87 was fixed at 9.15; all other KA values and all T0,variant values were jointly optimized. Tmax = 13.61 RPU and operator number = 7 were fixed globally.", False),
        ],
    )


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = set_run_font(paragraph.add_run("Mammalian fitted parameter tables"), size=8.5)
    run.font.color.rgb = RGBColor(105, 105, 105)

    document.core_properties.title = "Mammalian fitted DBD and LBD parameters"
    document.core_properties.subject = (
        "Current constrained 25-sensor mammalian CIC fit; LBD and DBD parameter tables"
    )
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def validate_inputs(
    lbd_order: list[str],
    lbd_rows: dict[str, dict[str, str]],
    dbd_order: list[str],
    dbd_rows: dict[str, dict[str, str]],
) -> None:
    if set(lbd_order) != set(LBD_ORDER):
        raise ValueError(
            f"Unexpected LBD groups. Output={sorted(lbd_order)}; expected={sorted(LBD_ORDER)}"
        )
    if len(dbd_order) != 11:
        raise ValueError(f"Expected 11 DBD parameter groups; found {len(dbd_order)}")
    if lbd_rows.keys() != set(LBD_ORDER):
        raise ValueError("LBD output index mismatch")
    if dbd_rows["lexAec87"]["KA_fixed"].casefold() != "true":
        raise ValueError("lexAec87 KA must be marked fixed")
    if not math.isclose(float(dbd_rows["lexAec87"]["KA"]), 9.15, rel_tol=0, abs_tol=1e-12):
        raise ValueError("lexAec87 KA must equal 9.15")
    if any(not math.isclose(float(row["Tmax"]), 13.61, abs_tol=1e-12) for row in dbd_rows.values()):
        raise ValueError("All DBD rows must carry fixed Tmax=13.61")
    if any(int(row["operator_number"]) != 7 for row in dbd_rows.values()):
        raise ValueError("All DBD rows must carry operator_number=7")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--result-dir", type=Path, default=FIT_DIR)
    parser.add_argument("--readme", type=Path, default=README)
    parser.add_argument("--output-docx", type=Path)
    args = parser.parse_args()
    output_docx = args.output_docx or args.result_dir / OUTPUT_DOCX.name
    lbd_order, lbd_rows = read_csv_indexed(args.result_dir / LBD_CSV.name, "canonical_lbd")
    dbd_order, dbd_rows = read_csv_indexed(args.result_dir / DBD_CSV.name, "canonical_dbd")
    dbd_aliases = read_dbd_aliases(args.readme)
    validate_inputs(lbd_order, lbd_rows, dbd_order, dbd_rows)

    document = Document()
    configure_document(document)
    build_lbd_table(document, lbd_rows)
    build_dbd_table(document, dbd_order, dbd_rows, dbd_aliases)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    print(output_docx)


if __name__ == "__main__":
    main()
